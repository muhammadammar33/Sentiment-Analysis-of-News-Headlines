# Converts docs/step_by_step_guide.md into docs/step_by_step_guide.docx.
# Run from the repository root: python scripts/build_guide_docx.py
import datetime
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)

def add_runs(p, text):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*"):
            p.add_run(part[1:-1]).italic = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10)
        else:
            p.add_run(part)

def shade(paragraph, colour="F2F2F0"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), colour)
    pPr.append(shd)

lines = open("docs/step_by_step_guide.md", encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("```"):
        i += 1
        block = []
        while not lines[i].startswith("```"):
            block.append(lines[i]); i += 1
        p = doc.add_paragraph(); shade(p)
        p.paragraph_format.left_indent = Inches(0.2)
        r = p.add_run("\n".join(block)); r.font.name = "Consolas"; r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    elif ln.startswith("# "):
        doc.add_heading(ln[2:], level=0)
    elif ln.startswith("## "):
        doc.add_heading(ln[3:], level=1)
    elif ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.match(r"^\|[-:| ]+\|$", lines[i]):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = t.cell(r_i, c_i); cell.text = ""
                add_runs(cell.paragraphs[0], val)
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(9.5)
        doc.add_paragraph()
        continue
    elif re.match(r"^\d+\. ", ln):
        add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", ln))
    elif ln.startswith("- "):
        add_runs(doc.add_paragraph(style="List Bullet"), ln[2:])
    elif ln.strip():
        add_runs(doc.add_paragraph(), ln)
    i += 1

cp = doc.core_properties
cp.author = cp.last_modified_by = "Muhammad Ammar"
cp.title = "Step-by-Step Guide - Sentiment Analysis of News Headlines"
cp.comments = ""
cp.created = cp.modified = datetime.datetime(2026, 10, 24)
doc.save("docs/step_by_step_guide.docx")
