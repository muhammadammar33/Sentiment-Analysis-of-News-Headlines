# Builds report/report.docx (from report/report.md) and presentation/presentation.pptx.
# Run from the repository root after the notebook has produced outputs/figures/.
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN

FIG = "outputs/figures/"
md = open("report/report.md", encoding="utf-8").read()

# ---------------- Word report ----------------
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)

def add_runs(p, text):
    for part in re.split(r"(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("***"):
            r = p.add_run(part[3:-3]); r.bold = r.italic = True
        elif part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*"):
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            p.add_run(part)

figs_after = {"## 3. Results": None}
fig_inserts = {
    "**Finding 1": [(FIG + "fig1_sentiment_bar.png", "Figure 1. VADER sentiment labels (bar chart)"),
                    (FIG + "fig2_sentiment_pie.png", "Figure 2. Share of headlines by label (pie chart)")],
    "**Finding 2": [(FIG + "fig3_sentiment_by_source.png", "Figure 3. Sentiment profile of satire vs. real news")],
    "**Finding 3": [(FIG + "fig4_compound_distribution.png", "Figure 4. Distribution of compound scores")],
}

lines = md.splitlines()
i = 0
while i < len(lines):
    ln = lines[i]
    for key, figs in fig_inserts.items():
        if ln.startswith(key):
            for path, cap in figs:
                doc.add_picture(path, width=Inches(5.6 if "pie" not in path else 3.6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                c = doc.add_paragraph(); r = c.add_run(cap); r.italic = True; r.font.size = Pt(9)
                c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if ln.startswith("# "):
        doc.add_heading(ln[2:], level=0)
        p = doc.add_paragraph(); p.add_run("Mini Assignment – NLP Research Topic  |  Submission date: 26 October 2026").italic = True
    elif ln.startswith("## "):
        doc.add_heading(ln[3:], level=1)
    elif ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.match(r"^\|[-:| ]+\|$", lines[i]):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = t.cell(r_i, c_i); cell.text = ""
                add_runs(cell.paragraphs[0], val)
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(9)
        if len(rows[0]) == 5:
            t.autofit = False
            for col, w in zip(t.columns, [0.3, 3.3, 1.0, 0.8, 1.3]):
                col.width = Inches(w)
            for row in t.rows:
                for cell, w in zip(row.cells, [0.3, 3.3, 1.0, 0.8, 1.3]):
                    cell.width = Inches(w)
        doc.add_paragraph()
        continue
    elif re.match(r"^\d+\. ", ln):
        p = doc.add_paragraph(style="List Number"); add_runs(p, re.sub(r"^\d+\. ", "", ln))
    elif ln.startswith("- "):
        p = doc.add_paragraph(style="List Bullet"); add_runs(p, ln[2:])
    elif ln.strip():
        p = doc.add_paragraph(); add_runs(p, ln)
    i += 1

import datetime
now = datetime.datetime(2026, 10, 20)
for cp, t in [(doc.core_properties, "Sentiment Analysis of News Headlines - Report")]:
    cp.author = "Muhammad Ammar"; cp.title = t; cp.comments = ""; cp.last_modified_by = "Muhammad Ammar"; cp.created = now; cp.modified = now
doc.save("report/report.docx")

# ---------------- Presentation ----------------
INK = PRGB(0x0B, 0x0B, 0x0B); MUTED = PRGB(0x52, 0x51, 0x4E)
RED = PRGB(0xE3, 0x49, 0x48); BLUE = PRGB(0x2A, 0x78, 0xD6)
prs = Presentation(); prs.slide_width = PI(13.333); prs.slide_height = PI(7.5)
blank = prs.slide_layouts[6]

def title(s, text, sub=None):
    tb = s.shapes.add_textbox(PI(0.6), PI(0.4), PI(12.1), PI(1.0)).text_frame
    tb.word_wrap = True
    p = tb.paragraphs[0]; p.text = text; p.font.size = PP(30); p.font.bold = True; p.font.color.rgb = INK
    if sub:
        p2 = tb.add_paragraph(); p2.text = sub; p2.font.size = PP(16); p2.font.color.rgb = MUTED

def bullets(s, items, left=0.6, top=1.7, width=12.1, height=5.3, size=20):
    tf = s.shapes.add_textbox(PI(left), PI(top), PI(width), PI(height)).text_frame; tf.word_wrap = True
    for k, it in enumerate(items):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        sub = it.startswith("  ")
        p.text = ("– " if sub else "• ") + it.strip()
        p.font.size = PP(size - 3 if sub else size); p.font.color.rgb = MUTED if sub else INK
        p.level = 1 if sub else 0; p.space_after = PP(8)

def slide(t, sub=None):
    s = prs.slides.add_slide(blank); title(s, t, sub); return s

s = prs.slides.add_slide(blank)
tf = s.shapes.add_textbox(PI(0.8), PI(2.3), PI(11.7), PI(3)).text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.text = "Sentiment Analysis of News Headlines"; p.font.size = PP(44); p.font.bold = True; p.font.color.rgb = INK
p = tf.add_paragraph(); p.text = "Do computational sentiment labels match what headlines mean?"; p.font.size = PP(22); p.font.color.rgb = MUTED
p = tf.add_paragraph(); p.text = "Python · pandas · VADER  |  NLP Mini Assignment  |  26 October 2026"; p.font.size = PP(16); p.font.color.rgb = MUTED

s = slide("Research question & data")
bullets(s, ["Question: when VADER labels a headline positive / negative / neutral, does the label reflect its meaning?",
            "Dataset: News Headlines Dataset for Sarcasm Detection (Misra & Arora, 2023)",
            "  28,619 headlines: HuffPost (real news) + The Onion (satire), with an is_sarcastic flag",
            "  The satire flag gives an independent test of how VADER handles sarcasm",
            "Cleaning: 0 missing values; 116 duplicate headlines removed → 28,503 unique",
            "  e.g. \"'no way to prevent this,' says only nation where this regularly happens\" appears 12×"])

s = slide("Method: VADER + documented decision rules")
bullets(s, ["VADER: dictionary of ~7,500 human-rated words (−4 … +4) + rules for negation, intensifiers, 'but', punctuation, capitals",
            "Produces a compound score in [−1, +1]",
            "Decision rules (Hutto & Gilbert, 2014):",
            "  compound ≥ +0.05 → positive",
            "  compound ≤ −0.05 → negative",
            "  otherwise → neutral",
            "Note: the dataset is all lower case, so VADER's capitalisation rule never applies"])

s = slide("Results: a seemingly balanced picture")
s.shapes.add_picture(FIG + "fig1_sentiment_bar.png", PI(0.5), PI(1.5), height=PI(5.4))
bullets(s, ["Negative: 9,402 (33.0%)", "Neutral: 10,200 (35.8%)", "Positive: 8,901 (31.2%)",
            "Taken at face value: news is emotionally balanced", "…but is it?"], left=9.6, top=2.0, width=3.4, size=18)

s = slide("Finding 1: VADER cannot see satire")
s.shapes.add_picture(FIG + "fig3_sentiment_by_source.png", PI(1.1), PI(1.5), width=PI(11))
bullets(s, ["Satire and real news get almost the same label mix; mean compound −0.016 vs −0.011",
            "χ² is 'significant' only because n is huge; Cramér's V ≈ 0.03 → negligible effect"], top=5.75, height=1.6, size=18)

s = slide("Finding 2: 'Neutral' often means 'unknown'")
s.shapes.add_picture(FIG + "fig4_compound_distribution.png", PI(0.4), PI(1.6), width=PI(7.6))
bullets(s, ["95.6% of neutral headlines contain no VADER dictionary word at all",
            "Missing: dies, shooting, hurricane, flood, layoffs, deported",
            "  'died' is in the dictionary; 'dies' is not",
            "  \"cilla black dies at 72\" → 0.00",
            "Finding 3: widening the neutral band to ±0.20 makes 43% neutral; ±0.50 makes 75%"],
        left=8.2, top=1.8, width=4.8, size=17)

s = slide("Close reading: 10 headlines (VADER agrees on 2/10)")
rows = [("Headline", "VADER", "Human", "Why"),
        ("student killed herself after university mishandled her rape report", "neg −0.88", "neg", "explicit cue ✓"),
        ("mexico mayor killed less than a day after taking office", "neg −0.67", "neg", "explicit cue ✓"),
        ("british singer … cilla black dies at 72", "neu 0.00", "neg", "missing word"),
        ("how fake meat may change our future thanksgiving dinners", "neg −0.48", "neu", "word sense"),
        ("duke students refuse to read 'fun home' …", "pos +0.27", "neg", "book title"),
        ("bob hope happy to see so many troops in heaven", "pos +0.87", "neg", "name + irony"),
        ("…underfunded schools leave students perfectly prepared…", "pos +0.82", "neg", "sarcasm"),
        ("'no way to prevent this,' says only nation …", "neu +0.03", "neg", "satire / implicit"),
        ("boeing lays off only guy who knows how to keep wings on plane", "neu 0.00", "neg", "implicit"),
        ("separating from a wonderful man after 27 years", "pos +0.57", "neg/mixed", "mixed")]
tbl = s.shapes.add_table(len(rows), 4, PI(0.5), PI(1.5), PI(12.3), PI(5.6)).table
for c, w in enumerate([7.3, 1.6, 1.4, 2.0]):
    tbl.columns[c].width = PI(w)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c); cell.text = val
        para = cell.text_frame.paragraphs[0]; para.font.size = PP(13); para.font.bold = (r == 0)
        if r and c == 1:
            para.font.color.rgb = RED if val.startswith("neg") else BLUE if val.startswith("pos") else MUTED

s = slide("Why VADER goes wrong")
bullets(s, ["Sarcasm & irony: positive words carry negative meaning (\"perfectly prepared\")",
            "Ambiguity: word sense (fake meat) and names (Bob Hope, Fun Home) are scored as emotions",
            "Context & world knowledge: meaning lies in the situation, not in emotional words (a plane without wings)",
            "Dictionary coverage: missing words → false 'neutral', so negative news is hidden",
            "What is being measured? topic vs. writer's stance vs. reader's reaction are mixed together",
            "Thresholds & preprocessing: category sizes are partly a researcher's choice"])

s = slide("Conclusion")
bullets(s, ["VADER is fast and transparent, and works when emotion is stated directly",
            "Its balanced-looking output (33 / 36 / 31%) hides three problems:",
            "  it does not tell satire from sincere reporting",
            "  most 'neutral' labels come from missing vocabulary",
            "  category sizes depend on arbitrary thresholds",
            "Use computational labels as a first rough pass that must be checked by reading the headlines",
            "Next steps: human-annotated random sample (with inter-annotator agreement); compare with a transformer model"])

s = slide("References")
bullets(s, ["Hutto, C. J., & Gilbert, E. (2014). VADER: A parsimonious rule-based model for sentiment analysis of social media text. ICWSM, 8(1), 216–225.",
            "Misra, R., & Arora, P. (2023). Sarcasm detection using news headlines dataset. AI Open, 4, 13–18.",
            "Misra, R., & Grover, J. (2021). Sculpting Data for ML: The first act of Machine Learning.",
            "Dataset: github.com/rishabhmisra/News-Headlines-Dataset-For-Sarcasm-Detection"], size=16)

cp = prs.core_properties
cp.author = "Muhammad Ammar"; cp.title = "Sentiment Analysis of News Headlines - Presentation"; cp.comments = ""
cp.last_modified_by = "Muhammad Ammar"; cp.created = now; cp.modified = now
prs.save("presentation/presentation.pptx")
